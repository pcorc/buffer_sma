"""Readable batch reports: cohort-aligned curves, downside rankings and run links."""
from dataclasses import asdict
from html import escape
import json
from pathlib import Path

import pandas as pd

from backtesting.research import comparison_table
from config.identity import rule_fingerprint
from config.settings import ScoreConfig
from studio.service import describe_rule


def export_batch_comparison(results, cases, folders, output, *, plots="top5", chart_cohort=None):
    """Results contain summary/daily/parameters only, keeping batch memory bounded."""
    output = Path(output)
    if plots not in ('top5', 'none'):
        raise ValueError('Plot mode must be top5 or none.')
    table = comparison_table(results)
    table['Mean Annualized Return'] = [r['summary']['Annualized Return'].mean() for r in results.values()]
    table['Total Return Rank'] = table['Mean Total Return'].rank(ascending=False, method='min', na_option='bottom').astype(int)
    table['Sharpe Rank'] = table['Mean Sharpe'].rank(ascending=False, method='min', na_option='bottom').astype(int)
    table['Drawdown Rank'] = table['Worst Drawdown'].rank(ascending=False, method='min', na_option='bottom').astype(int)
    table['Sortino Rank'] = table['Mean Sortino'].rank(ascending=False, method='min', na_option='bottom').astype(int)
    table['Turnover Rank'] = table['Mean Annualized Turnover'].rank(ascending=True, method='min', na_option='bottom').astype(int)
    configs = []
    for i, case in enumerate(cases, 1):
        e = case.experiment
        cfg = e.backtest
        config = {'Experiment': e.name, 'Configuration ID': f'C{i:03d}', 'Universe': ', '.join(cfg.universe), 'N': cfg.portfolio_size,
                  'Review': cfg.review, 'Trigger': 'ecr' if cfg.trigger == 'ecr' else cfg.buffer_metric,
                  'Trigger Rule': 'drop' if cfg.trigger == 'buffer_drop' else 'floor',
                  'Threshold': {'ecr': cfg.threshold, 'buffer_drop': cfg.buffer_drop}.get(cfg.trigger, cfg.metric_floor),
                  'Rule': describe_rule(cfg), 'Selection': cfg.selection, 'Execution Lag': cfg.execution_lag,
                  'Cost bps': cfg.cost_bps, 'Comparator': cfg.comparator,
                  'Score Profile': 'Standard ECR' if e.score == ScoreConfig() else json.dumps(asdict(e.score)),
                  'Rule Fingerprint': rule_fingerprint(e.score, cfg),
                  'Score improvement': cfg.min_selection_improvement, 'Max replacements': cfg.max_replacements_per_review,
                  'Turnover budget': cfg.max_turnover_per_review, 'Weight band': cfg.weight_band,
                  'Report folder': folders[e.name].relative_to(output).as_posix()}
        configs.append(config)
    config_table = pd.DataFrame(configs).set_index('Experiment')
    table = table.join(config_table.drop(columns=['N']))
    table = table.sort_values(['Drawdown Rank', 'Sortino Rank', 'Turnover Rank'], kind='stable')
    cohorts = pd.concat([r['summary'].assign(Experiment=name) for name,r in results.items()], ignore_index=True)
    table.to_csv(output / 'comparison.csv')
    cohorts.to_csv(output / 'cohort_comparison.csv', index=False)
    config_table.to_csv(output / 'configurations.csv')
    launches = next(iter(results.values()))['summary'].Cohort.tolist()
    chart_cohort = chart_cohort.title() if chart_cohort else launches[0]
    if chart_cohort not in launches:
        raise ValueError('Chart cohort must be one of the simulated launch months.')
    winners = top_five_tables(table)
    shortlisted = pd.concat([rows.assign(Criterion=RANKINGS[key][0], Rank=range(1,len(rows)+1)) for key,rows in winners.items()])
    shortlisted.to_csv(output / 'top_five.csv')
    curves = {}
    with pd.ExcelWriter(output / 'comparison.xlsx', engine='openpyxl') as writer:
        table.to_excel(writer, sheet_name='Comparison')
        shortlisted.to_excel(writer, sheet_name='Top Five')
        cohorts.to_excel(writer, sheet_name='Launch Cohorts', index=False)
        config_table.to_excel(writer, sheet_name='Configurations')
        pd.DataFrame({'Read me': [
            'Sorted by least severe worst-cohort drawdown, then Sortino and turnover. This is a research ordering, not a validated optimum.',
            'Each launch month is a separate backtest. Aggregate columns are arithmetic means of those tests, not returns of one investable portfolio.',
            'Top Five ranks mean total return, least severe worst-cohort drawdown, and mean Sharpe across all launches. Nonfinite scores are excluded; ties use experiment name.',
            'Compare curves within the same launch month. Windows must match exactly; missing sessions are not filled.',
            'All scenarios select highest multiplicative ECR, use next-close execution, and share the specified transaction cost.',
            'Inspect the individual report folders for decisions, holdings, candidates, and diagnostics.',
            'Thresholds, portfolio size and turnover rules are in Configurations; full replay settings are in each run/config.json.',
        ]}).to_excel(writer, sheet_name='Read Me', index=False)
        for cohort in launches:
            series = {name: result['daily'].loc[result['daily'].Cohort == cohort].set_index('Date')['Strategy NAV']
                      for name,result in results.items()}
            frame = pd.concat(series, axis=1)
            # Never compare a backfilled or silently shortened price series.
            if frame.isna().any().any():
                raise ValueError(f'{cohort}: scenario NAV calendars differ or contain missing values.')
            curves[cohort] = frame
            frame.to_excel(writer, sheet_name=f'NAV {cohort}')
        from openpyxl.styles import Font, PatternFill
        for sheet in writer.book:
            sheet.freeze_panes = 'B2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='1B3442')
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(60, max(16, len(str(column[0].value or ''))+2))
            # Return, drawdown and turnover columns remain numeric and easy to compare.
            for cell in sheet[1]:
                label = str(cell.value or '')
                if any(word in label for word in ('Return', 'Drawdown', 'Deviation', 'Turnover', 'Cumulative Excess', 'vs BUFR', 'Weight band')) and 'Rank' not in label:
                    for row in sheet.iter_rows(min_row=2, min_col=cell.column, max_col=cell.column):
                        row[0].number_format = '0.00%'

    plot_links = []
    if plots == 'top5':
        plot_links = render_top_five(curves, winners, cases, output, chart_cohort)
    links = ''.join(f'<li><a href="{url}">{escape(label)}</a></li>' for label,url in plot_links)
    displays = ''.join(f'<h2>{escape(RANKINGS[key][0])}</h2>' + rows[
        ['Configuration ID','Portfolio Size','Mean Total Return','Worst Drawdown','Mean Sharpe','Review','Trigger','Threshold',
         'Score improvement','Max replacements','Turnover budget','Weight band']
    ].to_html(float_format=lambda v:f'{v:.4f}',escape=True) for key,rows in winners.items())
    runs = ''.join(f'<tr><td>{escape(name)}</td><td><a href="{folder.relative_to(output).as_posix()}/ecr_results.xlsx">Excel</a></td></tr>' for name,folder in folders.items())
    chart_section = ('<h2>Top-five comparison charts</h2><p>Each chart includes the parameter sets and matching line IDs. HTML charts switch between launch months; PNGs show '
                     + escape(chart_cohort) + '.</p><ul>' + links + '</ul>') if plots == 'top5' else '<p>All plotting was disabled for this batch.</p>'
    (output/'index.html').write_text(f'''<!doctype html><html lang="en"><meta charset="utf-8"><title>Conservative ECR comparison</title>
<style>body{{font:16px Arial;color:#1b3442;max-width:1600px;margin:40px auto;padding:0 24px}}h1{{background:#f9eed7;padding:24px;border-left:4px solid #eaaa43}}table{{border-collapse:collapse;width:100%;font-size:13px}}td,th{{padding:10px;border-bottom:1px solid #ddd;text-align:left}}a{{color:#1672c4}}</style>
<h1>Conservative ECR · {len(results)} configurations</h1><p>Highest ECR selects every replacement. Individual iteration charts are suppressed; all Excel/CSV data and replay settings remain available.</p>
<p><a href="comparison.xlsx">Comparison workbook</a> · <a href="comparison.csv">Comparison CSV</a> · <a href="top_five.csv">Top-five parameter sets</a></p>
{chart_section}<p>Rankings use mean total return, least severe worst-cohort drawdown, and mean Sharpe across all launches. These are research summaries, not a combined portfolio or an out-of-sample recommendation. Nonfinite ranking values are excluded; ties use experiment name. Returns/drawdowns below are decimals (0.10 = 10%).</p>{displays}
<h2>Individual reports</h2><table><tr><th>Configuration</th><th>Tables</th></tr>{runs}</table></html>''',encoding='utf-8')
    return table


RANKINGS = {
    'total-return': ('Top five · highest total return', 'Mean Total Return', 'Mean total return'),
    'drawdown': ('Top five · smallest maximum drawdown', 'Worst Drawdown', 'Worst drawdown'),
    'sharpe': ('Top five · highest Sharpe', 'Mean Sharpe', 'Mean Sharpe'),
}


def top_five_tables(table):
    """Independent rankings; negative drawdowns closer to zero are better."""
    import numpy as np
    return {key: table.loc[np.isfinite(table[column])].sort_index().sort_values(column, ascending=False, kind='stable').head(5)
            for key, (_, column, _) in RANKINGS.items()}


def render_top_five(curves, winners, cases, output, chart_cohort):
    """Only three chart pairs, with cohort controls and parameter tables inside."""
    import textwrap
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import plotly.graph_objects as go
    from plotly.subplots import make_subplots

    charts = output / 'charts'
    charts.mkdir()
    colors = ['#1B3442', '#D08A18', '#1672C4', '#7B4F98', '#25856A']
    case_map = {c.experiment.name: c for c in cases}
    ids = {c.experiment.name: f'C{i:03d}' for i,c in enumerate(cases,1)}
    links = []
    for key, selected in winners.items():
        title, metric, metric_label = RANKINGS[key]
        if selected.empty:
            continue
        names = selected.index.tolist()
        universes = list(dict.fromkeys(selected.Universe))
        universe_ids = {value: f'U{i:02d}' for i,value in enumerate(universes,1)}
        headers = ['ID', 'N', 'Trigger', 'Review', 'ECR gain', 'Max swaps', 'Budget', 'Band (pp)', 'Universe', metric_label]
        rows = []
        for name, row in selected.iterrows():
            trigger = {'ecr':'ECR','buffer_integrity':'Buffer','cap_integrity':'Cap','dbb':'DBB'}[row.Trigger]
            rule = f'{trigger} drop {row.Threshold:.0%}' if row['Trigger Rule'] == 'drop' else f'{trigger} < {row.Threshold:g}'
            rows.append([ids[name], str(int(row['Portfolio Size'])), rule, row.Review,
                         f'{row["Score improvement"]:g}', str(int(row['Max replacements'])) if row['Max replacements'] else 'All',
                         f'{row["Turnover budget"]:.0%}', 'Drift' if pd.isna(row['Weight band']) else f'±{row["Weight band"]*100:g}',
                         universe_ids[row.Universe], f'{row[metric]:.3f}' if key=='sharpe' else f'{row[metric]:.2%}'])
        notes = []
        for value, uid in universe_ids.items():
            example = next(case_map[name] for name in names if selected.loc[name,'Universe']==value)
            score = example.experiment.score
            notes.append(f'{uid}: {value}. Period center/width {score.period_center:g}/{score.period_width:g} days.')
        common = case_map[names[0]].experiment.backtest
        notes.append(f'Common: highest ECR selection; next-close execution; {common.cost_bps:g} bps. ECR gain = minimum replacement improvement; max swaps and budget apply per review.')
        notes.append('Ranked across all launch tests; each curve shows one launch. Band is absolute percentage points around 1/N; Drift means no weight reset.')
        wrapped_notes = '\n'.join(textwrap.fill(note, 155) for note in notes)

        fig = make_subplots(rows=3,cols=1,specs=[[{'type':'xy'}],[{'type':'xy'}],[{'type':'table'}]],
                            row_heights=[.42,.32,.26],vertical_spacing=.1,
                            subplot_titles=['Growth of 100','Drawdown','Parameter sets · IDs match the lines'])
        trace_names = []
        for i,name in enumerate(names):
            frame = curves[chart_cohort][name]
            label = f'{ids[name]} · N={int(selected.loc[name,"Portfolio Size"])} · {rows[i][2]} · {rows[i][3]}'
            for row, values in [(1,frame),(2,frame/frame.cummax()-1)]:
                fig.add_trace(go.Scatter(x=frame.index,y=values,name=label,legendgroup=name,
                                        showlegend=row==1,line=dict(color=colors[i],width=2)),row=row,col=1)
                trace_names.append((name,row))
        fig.add_trace(go.Table(header=dict(values=headers,fill_color='#1B3442',font=dict(color='white'),align='left'),
                               cells=dict(values=list(zip(*rows)),fill_color='#F5F4EF',align='left',height=30)),row=3,col=1)
        buttons = []
        for cohort, frame in curves.items():
            xs, ys = [], []
            for name,row in trace_names:
                values = frame[name]
                xs.append(frame.index.tolist())
                ys.append((values if row==1 else values/values.cummax()-1).tolist())
            buttons.append(dict(label=f'{cohort} launch',method='update',args=[{'x':xs,'y':ys},
                           {'title':{'text':f'{title} · {cohort} launch'}} ,list(range(len(trace_names)))]))
        fig.update_layout(template='plotly_white',height=1120,title=f'{title} · {chart_cohort} launch',
                          font=dict(family='Arial',color='#1B3442'),margin=dict(l=55,r=35,t=130,b=45),
                          legend=dict(orientation='h',y=1.1,groupclick='togglegroup'),
                          updatemenus=[dict(buttons=buttons,active=list(curves).index(chart_cohort),x=1,y=1.2,xanchor='right')],
                          hovermode='x unified')
        fig.update_yaxes(tickformat='.0%',row=2,col=1)
        html_path = charts/f'top5-{key}.html'
        fig.write_html(html_path,include_plotlyjs='directory')
        # Explanatory notes travel with the interactive chart, not just the index.
        html = html_path.read_text(encoding='utf-8')
        html_path.write_text(html.replace('</body>', '<div style="font:14px Arial;color:#1B3442;padding:0 55px 30px;white-space:pre-wrap">'+escape(wrapped_notes)+'</div></body>'),encoding='utf-8')

        figure = plt.figure(figsize=(16,12))
        grid = figure.add_gridspec(3,1,height_ratios=[3,2,1.6],hspace=.4)
        axes = [figure.add_subplot(grid[i]) for i in range(3)]
        for i,name in enumerate(names):
            values = curves[chart_cohort][name]
            axes[0].plot(values.index,values,label=ids[name],color=colors[i],linewidth=1.8)
            axes[1].plot(values.index,(values/values.cummax()-1)*100,color=colors[i],linewidth=1.8)
        axes[0].set_title('Growth of 100',loc='left');axes[0].legend(ncol=5,loc='upper left')
        axes[1].set_title('Drawdown (%)',loc='left')
        for ax in axes[:2]:
            ax.grid(alpha=.15)
            ax.spines[['top','right']].set_visible(False)
        axes[2].axis('off')
        tab = axes[2].table(cellText=rows,colLabels=headers,cellLoc='left',loc='upper center',colWidths=[.06,.03,.13,.09,.07,.07,.07,.08,.07,.12])
        tab.auto_set_font_size(False);tab.set_fontsize(9);tab.scale(1,1.6)
        for (row,col),cell in tab.get_celld().items():
            cell.set_edgecolor('#DFE3E5')
            if row==0:
                cell.set_facecolor('#1B3442');cell.get_text().set_color('white')
            elif col==0:
                cell.set_facecolor(colors[row-1]);cell.get_text().set_color('white')
        figure.suptitle(f'{title} · {chart_cohort} launch',fontsize=17,color='#1B3442')
        figure.text(.07,.02,wrapped_notes,fontsize=8,va='bottom',color='#536C7A')
        figure.subplots_adjust(top=.93,bottom=.17)
        png_path = charts/f'top5-{key}.png'
        figure.savefig(png_path,dpi=140,bbox_inches='tight');plt.close(figure)
        links.extend([(title+' · interactive',f'charts/{html_path.name}'),(title+f' · {chart_cohort} PNG',f'charts/{png_path.name}')])
    return links
