"""Audit a comparison workbook and evaluate simple portfolios on matched stress dates."""
import argparse
from dataclasses import asdict
import json
from pathlib import Path
import sys

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import pandas as pd
from backtesting.regimes import REGIMES, load_export, evaluate_windows, fresh_checks, drawdown_heatmap
from backtesting.conservative_runner import ensure_archive_root
from backtesting.data_pipeline import file_hash
from studio.archive import DEFAULT_ARCHIVE_ROOT, new_archive_path

NOTES = [
    'Carry-in analysis evaluates the already-invested portfolios in the uploaded NAV sheets. It does not restart holdings at the stress boundary.',
    'All cases within each regime use the exact same starting and ending closing observations. No missing boundary is filled; unavailable windows are listed in Coverage.',
    'Returns run from the start close to the end close. Drawdown resets its high-water mark at the window start; it is not the full-history underwater drawdown.',
    'Regime Summary shows mean returns and the worst drawdown across original launch cohorts. Those cohorts overlap and are not independent market-regime trials.',
    'Fresh Start Results, when requested, use new highest-ECR portfolios before the stress window. They are separate tests, never pooled with carry-in results.',
    'Simple cases have no ECR improvement hurdle, no replacement-count limit, no restrictive turnover budget and no weight-band reset. Cadence defaults to monthly.',
    'The fresh-start engine uses current repository histories, default multiplicative ECR, strict triggers, next-close execution and zero trading costs. Exact replay settings and data hashes are in fresh_runs.',
    'The 2024 fresh start is May 20, not June 24: the current benchmark file lacks both SPY and BUFR values at the June launch. No benchmark values are invented.',
    'Fresh simulations extend beyond stress boundaries to the declared simulation end; reported stress performance stops at the exact regime end. Simulation is chronological.',
    'Sharpe reconciliation uses the engine convention of a zero risk-free rate and sqrt(252). Short stress windows are evaluated by loss and drawdown, not optimized on annualized Sharpe.',
    'Stress dates are retrospective diagnostics. No claim of a validated optimal threshold follows from these selected periods. Recovery/expansion windows supply context.',
    'The 2020 COVID window cannot be evaluated from a workbook starting in 2021. The present 12-fund universe was not all available in early 2020.',
]


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook',type=Path,required=True,help='Path to an exported comparison.xlsx with NAV launch sheets.')
    parser.add_argument('--archive-root',default=DEFAULT_ARCHIVE_ROOT)
    parser.add_argument('--review',choices=['monthly','quarterly','all'],default='monthly')
    parser.add_argument('--fresh-starts',action='store_true',help='Also simulate fresh portfolios before four stress periods using repository data.')
    parser.add_argument('--no-plots',action='store_true',help='Omit the one or two drawdown heatmaps.')
    args=parser.parse_args(argv)
    folder=None
    try:
        root=ensure_archive_root(args.archive_root)
        print('Auditing workbook summaries against NAV histories...',flush=True)
        configs,curves,audit,original=load_export(args.workbook,args.review)
        details,summary,coverage=evaluate_windows(curves)
        print(f'{len(audit)} summary rows reconciled; {len(configs)} simple configurations selected.',flush=True)
        folder=new_archive_path(root,'Conservative-regimes','analysis')
        folder.mkdir(parents=True,exist_ok=False)
        fresh,plans=pd.DataFrame(),pd.DataFrame()
        sources={}
        if args.fresh_starts:
            fresh,plans,sources=fresh_checks(configs,folder)
        tables={'Simple Configurations':configs,'Original Summary':original.loc[original.Experiment.isin(configs.Experiment)],
                'Regime Summary':summary,'Carry-in Detail':details,'Coverage':coverage,'NAV Audit':audit,
                'Window Definitions':pd.DataFrame([asdict(r) for r in REGIMES]),
                'Read Me':pd.DataFrame({'Methodology':NOTES})}
        if not fresh.empty:
            tables['Fresh Start Results']=fresh
            tables['Fresh Start Plan']=plans
        with pd.ExcelWriter(folder/'regime_analysis.xlsx',engine='openpyxl') as writer:
            from openpyxl.styles import Font, PatternFill
            for name,frame in tables.items():
                frame.to_excel(writer,sheet_name=name,index=False)
                frame.to_csv(folder/(name.lower().replace(' ','_')+'.csv'),index=False)
                sheet=writer.book[name];sheet.freeze_panes='C2';sheet.auto_filter.ref=sheet.dimensions
                for cell in sheet[1]:
                    cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='1B3442')
                    sheet.column_dimensions[cell.column_letter].width=22 if cell.value!='Experiment' else 75
                    if ('Return' in str(cell.value) or 'Drawdown' in str(cell.value)) and 'Error' not in str(cell.value):
                        for row in sheet.iter_rows(min_row=2,min_col=cell.column,max_col=cell.column):row[0].number_format='0.00%'
        manifest={'source_workbook':str(args.workbook.resolve()),'source_sha256':file_hash(args.workbook),
                  'review':args.review,'fresh_starts':args.fresh_starts,'sources':sources,
                  'regimes':[asdict(r) for r in REGIMES],'simple_configurations':len(configs)}
        (folder/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
        (folder/'README.txt').write_text('\n\n'.join(NOTES)+'\n')
        if not args.no_plots:
            drawdown_heatmap(summary,configs,folder)
            if not fresh.empty:
                drawdown_heatmap(fresh,configs,folder,filename='fresh-stress-drawdowns.png',value='Max Drawdown')
        (folder/'ARCHIVE_COMPLETE.txt').write_text('Regime evaluation complete. Open regime_analysis.xlsx; read methodology before comparing carry-in and fresh starts.\n')
        print(f'Results: {folder / "regime_analysis.xlsx"}',flush=True)
        return folder
    except (ValueError,OSError,KeyError) as exc:
        parser.exit(1,f'Unable to complete regime analysis: {exc}\n'+(f'Partial output: {folder}\n' if folder else ''))


if __name__=='__main__':
    main()
